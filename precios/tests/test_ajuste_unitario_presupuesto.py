from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from precios.models import Cliente, Presupuesto, PresupuestoItem, PresupuestoItemAjusteUnitario
from precios.services.pdf_presupuesto import generar_pdf_presupuesto
from precios.services.excel_presupuesto import generar_excel_presupuesto


class AjusteUnitarioModeloTests(APITestCase):
    def setUp(self):
        self.cliente = Cliente.objects.create(
            numero=9008, nombre="Cliente Ajuste", incremento_dd=Decimal("0"), incremento_negra=Decimal("0"),
        )
        self.presupuesto = Presupuesto.objects.create(cliente=self.cliente, cotizacion_dolar=Decimal("1000"))
        self.item = PresupuestoItem.objects.create(
            presupuesto=self.presupuesto, tipo=PresupuestoItem.TIPO_OTRO, descripcion="Tubo 30x60",
            cantidad=Decimal("2"), precio_unitario_pesos=Decimal("1000"),
        )

    def test_sin_ajuste_el_precio_efectivo_es_el_original(self):
        self.assertEqual(self.item.precio_unitario_efectivo_pesos, Decimal("1000.00"))
        self.assertEqual(self.item.subtotal_pesos, Decimal("2000.00"))

    def test_ajuste_se_suma_al_precio_unitario_y_se_multiplica_por_cantidad(self):
        PresupuestoItemAjusteUnitario.objects.create(item=self.item, descripcion="Flete por unidad", monto_pesos=Decimal("50"))
        self.assertEqual(self.item.precio_unitario_efectivo_pesos, Decimal("1050.00"))
        # cantidad=2, entonces el ajuste de $50/u impacta $100 en el subtotal
        self.assertEqual(self.item.subtotal_pesos, Decimal("2100.00"))

    def test_varios_ajustes_se_acumulan(self):
        PresupuestoItemAjusteUnitario.objects.create(item=self.item, descripcion="A", monto_pesos=Decimal("50"))
        PresupuestoItemAjusteUnitario.objects.create(item=self.item, descripcion="B", monto_pesos=Decimal("30"))
        self.assertEqual(self.item.precio_unitario_efectivo_pesos, Decimal("1080.00"))

    def test_descripcion_es_opcional(self):
        ajuste = PresupuestoItemAjusteUnitario.objects.create(item=self.item, monto_pesos=Decimal("50"))
        self.assertEqual(ajuste.descripcion, "")

    def test_precio_unitario_original_no_cambia(self):
        PresupuestoItemAjusteUnitario.objects.create(item=self.item, monto_pesos=Decimal("50"))
        self.assertEqual(self.item.precio_unitario_pesos, Decimal("1000"))


class AjusteUnitarioAPITests(APITestCase):
    def setUp(self):
        self.usuario = get_user_model().objects.create_user(username="u1", password="x", is_staff=True)
        self.client.force_authenticate(self.usuario)

        self.cliente = Cliente.objects.create(
            numero=9009, nombre="Cliente Ajuste API", incremento_dd=Decimal("0"), incremento_negra=Decimal("0"),
        )
        self.presupuesto = Presupuesto.objects.create(cliente=self.cliente, cotizacion_dolar=Decimal("1000"))
        self.item = PresupuestoItem.objects.create(
            presupuesto=self.presupuesto, tipo=PresupuestoItem.TIPO_CHAPA, descripcion="Chapa DD · espesor 16",
            cantidad=Decimal("2"), precio_unitario_pesos=Decimal("1000"),
        )

    def test_agregar_ajuste_ok(self):
        resp = self.client.post(
            f"/api/presupuestos/{self.presupuesto.id}/items/{self.item.id}/agregar-ajuste-unitario/",
            {"descripcion": "Estampado", "monto_pesos": "50"},
        )
        self.assertEqual(resp.status_code, 200)
        self.item.refresh_from_db()
        self.assertEqual(self.item.ajustes_unitarios.count(), 1)
        self.assertEqual(self.item.precio_unitario_efectivo_pesos, Decimal("1050.00"))
        item_devuelto = resp.data["items"][0]
        self.assertEqual(item_devuelto["precio_unitario_efectivo_pesos"], "1050.00")
        self.assertEqual(item_devuelto["ajustes_unitarios"][0]["descripcion"], "Estampado")
        self.assertEqual(item_devuelto["descripcion_con_ajustes"], "Chapa Laminado en Frío · espesor 16 (Estampado)")

    def test_agregar_ajuste_sin_descripcion_falla(self):
        resp = self.client.post(
            f"/api/presupuestos/{self.presupuesto.id}/items/{self.item.id}/agregar-ajuste-unitario/",
            {"monto_pesos": "50"},
        )
        self.assertEqual(resp.status_code, 400)

    def test_agregar_ajuste_con_tipo_no_permitido_falla(self):
        resp = self.client.post(
            f"/api/presupuestos/{self.presupuesto.id}/items/{self.item.id}/agregar-ajuste-unitario/",
            {"descripcion": "Otro", "monto_pesos": "50"},
        )
        self.assertEqual(resp.status_code, 400)

    def test_agregar_ajuste_a_tubo_falla(self):
        tubo = PresupuestoItem.objects.create(
            presupuesto=self.presupuesto, tipo=PresupuestoItem.TIPO_TUBO, descripcion="Tubo 30x60",
            cantidad=Decimal("2"), precio_unitario_pesos=Decimal("1000"),
        )
        resp = self.client.post(
            f"/api/presupuestos/{self.presupuesto.id}/items/{tubo.id}/agregar-ajuste-unitario/",
            {"descripcion": "Estampado", "monto_pesos": "50"},
        )
        self.assertEqual(resp.status_code, 400)

    def test_agregar_ajuste_sin_monto_falla(self):
        resp = self.client.post(
            f"/api/presupuestos/{self.presupuesto.id}/items/{self.item.id}/agregar-ajuste-unitario/",
            {"descripcion": "Sin monto"},
        )
        self.assertEqual(resp.status_code, 400)

    def test_quitar_ajuste_ok(self):
        ajuste = PresupuestoItemAjusteUnitario.objects.create(item=self.item, monto_pesos=Decimal("50"))
        resp = self.client.delete(
            f"/api/presupuestos/{self.presupuesto.id}/items/{self.item.id}/ajustes-unitarios/{ajuste.id}/"
        )
        self.assertEqual(resp.status_code, 200)
        self.item.refresh_from_db()
        self.assertEqual(self.item.ajustes_unitarios.count(), 0)
        self.assertEqual(self.item.precio_unitario_efectivo_pesos, Decimal("1000.00"))


class AjusteUnitarioEnDocumentosTests(APITestCase):
    def setUp(self):
        self.cliente = Cliente.objects.create(
            numero=9010, nombre="Cliente Doc Ajuste", incremento_dd=Decimal("0"), incremento_negra=Decimal("0"),
        )
        self.presupuesto = Presupuesto.objects.create(
            cliente=self.cliente, cotizacion_dolar=Decimal("1000"), impuesto_pct=Decimal("21"),
        )
        self.item = PresupuestoItem.objects.create(
            presupuesto=self.presupuesto, tipo=PresupuestoItem.TIPO_CHAPA, descripcion="Chapa DD · espesor 16",
            cantidad=Decimal("2"), precio_unitario_pesos=Decimal("1000"),
        )
        PresupuestoItemAjusteUnitario.objects.create(item=self.item, descripcion="Estampado", monto_pesos=Decimal("50"))

    def test_pdf_con_ajuste_no_explota(self):
        pdf_bytes = generar_pdf_presupuesto(self.presupuesto)
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))

    def test_excel_con_ajuste_no_explota(self):
        excel_bytes = generar_excel_presupuesto(self.presupuesto)
        self.assertTrue(len(excel_bytes) > 0)
        ws = load_workbook(BytesIO(excel_bytes), data_only=True).active
        self.assertEqual(ws["A10"].value, "Chapa Laminado en Frío · espesor 16 (Estampado)")
        self.assertEqual(ws["D10"].value, "-")
        self.assertEqual(ws["E10"].value, 1050)
        self.assertEqual(ws["F10"].value, 2100)
        self.assertIsNone(ws["A11"].value)
