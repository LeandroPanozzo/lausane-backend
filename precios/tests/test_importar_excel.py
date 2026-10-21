import io
from decimal import Decimal
from unittest import mock

import openpyxl
from django.test import TestCase

from precios.models import Cliente, ConfiguracionGlobal, MarkupEspesorChapa, Tubo
from precios.services import importar_excel


def _construir_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INICIO"

    ws["C6"] = 1.3
    ws["C8"] = 1.1
    ws["A1"] = "% Incremento"
    ws["E1"] = 0.3

    ws["B3"] = "CLIENTES:"
    ws["B4"], ws["D4"], ws["F4"], ws["L4"] = 1, "Cliente Uno", 8, 5
    ws["B5"], ws["D5"], ws["F5"], ws["L5"] = 2, "Cliente Dos", 10, 6

    ws["A10"] = "Peso Teorico Tubo"
    ws["A11"], ws["B11"], ws["D11"], ws["E11"] = 1.2, 2.5, "30 50 0,9", 0.9

    ws_cliente = wb.create_sheet("1")
    ws_cliente["B6"] = "16"
    ws_cliente["C6"] = "=SUM(C16*0.07)+C16"
    ws_cliente["D6"] = "=SUM(D16*0.09)+D16"
    ws_cliente["E6"] = "=SUM(E16*0.11)+E16"
    ws_cliente["G6"] = "18"
    ws_cliente["H6"] = "=SUM(H16*0.04)+H16"
    ws_cliente["I6"] = "=SUM(I16*0.05)+I16"
    ws_cliente["J6"] = "=SUM(J16*0.06)+J16"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ImportarExcelTests(TestCase):
    def test_importa_configuracion_clientes_tubos_y_markups(self):
        resultado = importar_excel.importar_workbook(_construir_workbook())

        config = ConfiguracionGlobal.actual()
        self.assertEqual(config.precio_base_dd, Decimal("1.3"))
        self.assertEqual(config.precio_base_negra, Decimal("1.1"))
        self.assertEqual(config.incremento_tubos, Decimal("0.3"))

        self.assertEqual(Cliente.objects.count(), 2)
        uno = Cliente.objects.get(numero=1)
        self.assertEqual(uno.nombre, "Cliente Uno")
        self.assertEqual(uno.incremento_dd, Decimal("8"))

        self.assertEqual(Tubo.objects.count(), 1)
        self.assertEqual(Tubo.objects.get().dimensiones, "30 50 0,9")

        self.assertEqual(MarkupEspesorChapa.objects.count(), 6)
        self.assertEqual(resultado.clientes_creados, 2)
        self.assertEqual(resultado.tubos_creados, 1)
        self.assertEqual(resultado.markups_creados, 6)

    def test_import_es_atomico_si_falla_a_mitad_de_camino(self):
        """Si algo revienta a mitad del import (acá, forzado en tubos), no
        deben quedar datos a medio importar de los pasos anteriores."""
        with mock.patch.object(importar_excel, "importar_tubos", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                importar_excel.importar_workbook(_construir_workbook())

        self.assertEqual(Cliente.objects.count(), 0)
        self.assertFalse(ConfiguracionGlobal.objects.exists())
