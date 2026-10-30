"""Genera el Excel de un presupuesto (misma info que el PDF, formato .xlsx)."""

from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from precios.models import Presupuesto
from precios.services.texto_presupuesto import descripcion_item_con_ajustes

ROJO = "E30613"
GRIS_CLARO = "F4F4F5"


def generar_excel_presupuesto(presupuesto: Presupuesto) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"

    ws["A1"] = "PRESUPUESTO LAUSANE"
    ws["A1"].font = Font(size=16, bold=True, color=ROJO)
    ws.merge_cells("A1:F1")

    ws["A2"] = f"N° {presupuesto.codigo}"
    ws["A2"].font = Font(size=11, bold=True)
    ws.merge_cells("A2:F2")

    fecha_local = timezone.localtime(presupuesto.creado).strftime("%d/%m/%Y")
    ws["A4"] = "Cliente:"
    ws["B4"] = presupuesto.cliente.nombre
    ws["A5"] = "N° de cliente:"
    ws["B5"] = presupuesto.cliente.numero
    ws["A6"] = "Fecha:"
    ws["B6"] = fecha_local
    ws["A7"] = "Cotización dólar aplicada:"
    ws["B7"] = float(presupuesto.cotizacion_dolar)
    for fila in range(4, 8):
        ws[f"A{fila}"].font = Font(bold=True)

    fila_header = 9
    encabezados = ["Descripción", "Cantidad", "Unidad", "Longitud (mm)", "Precio unitario ($)", "Subtotal ($)"]
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_header, column=col, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=ROJO)
        celda.alignment = Alignment(horizontal="center")

    fila = fila_header + 1
    for i, item in enumerate(presupuesto.items.all()):
        ws.cell(row=fila, column=1, value=descripcion_item_con_ajustes(item))
        ws.cell(row=fila, column=2, value=float(item.cantidad)).alignment = Alignment(horizontal="right")
        ws.cell(row=fila, column=2).number_format = "0.###"
        ws.cell(row=fila, column=3, value=item.unidad).alignment = Alignment(horizontal="center")
        celda_longitud = ws.cell(row=fila, column=4, value=item.longitud_mm if item.longitud_mm else "-")
        celda_longitud.alignment = Alignment(horizontal="center")
        ws.cell(row=fila, column=5, value=float(item.precio_unitario_efectivo_pesos)).number_format = "#,##0.00"
        ws.cell(row=fila, column=6, value=float(item.subtotal_pesos)).number_format = "#,##0.00"
        if i % 2 == 1:
            for col in range(1, 7):
                ws.cell(row=fila, column=col).fill = PatternFill("solid", fgColor=GRIS_CLARO)
        fila += 1
        for extra in item.extras.all():
            celda_desc = ws.cell(row=fila, column=1, value=f"   + {extra.descripcion}")
            celda_desc.font = Font(italic=True, size=9, color="666666")
            ws.cell(row=fila, column=6, value=float(extra.precio_pesos)).number_format = "#,##0.00"
            fila += 1

    fila += 1
    ws.cell(row=fila, column=5, value="TOTAL").font = Font(bold=True)
    ws.cell(row=fila, column=5).alignment = Alignment(horizontal="right")
    total_celda = ws.cell(row=fila, column=6, value=float(presupuesto.total_pesos))
    total_celda.font = Font(bold=True, size=12, color=ROJO)
    total_celda.number_format = "#,##0.00"

    if presupuesto.notas:
        fila += 2
        ws.cell(row=fila, column=1, value=f"Notas: {presupuesto.notas}").font = Font(italic=True, size=9)

    anchos = [46, 12, 12, 14, 18, 18]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
